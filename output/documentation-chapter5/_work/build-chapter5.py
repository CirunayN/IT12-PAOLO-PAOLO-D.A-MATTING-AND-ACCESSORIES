from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import hashlib
import json
import re
import textwrap

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'output/documentation-chapter5'
SOURCE = Path(r'C:/Users/charl/OneDrive/Desktop/Internet/Downloads/Documentation_IT12_Chapter 1-5.docx')
FINAL = OUT / 'Documentation_IT12_Chapter_1-5_Completed.docx'
PART = OUT / '_work/chapter5-appendices.docx'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.5), Inches(11)
s.top_margin = s.bottom_margin = Inches(1)
s.left_margin, s.right_margin = Inches(1.5), Inches(1)

def style(name, font, size, before=0, after=0, line=1.0, bold=False):
    x = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    x.font.name = font
    x.font.size = Pt(size)
    x.font.color.rgb = RGBColor(0, 0, 0)
    x.font.bold = bold
    x.paragraph_format.space_before = Pt(before)
    x.paragraph_format.space_after = Pt(after)
    x.paragraph_format.line_spacing = line
    x.paragraph_format.widow_control = True
    return x

style('Ch5Body', 'Times New Roman', 12, after=8, line=1.5)
style('Ch5Heading', 'Times New Roman', 14, before=6, after=10, bold=True).paragraph_format.keep_with_next = True
style('Ch5Subheading', 'Times New Roman', 12, before=8, after=6, bold=True).paragraph_format.keep_with_next = True
style('Ch5Source', 'Times New Roman', 10, after=7, line=1.1).paragraph_format.keep_with_next = True
style('Ch5Code', 'Courier New', 9.5, line=1.0)
d.styles['Ch5Code'].paragraph_format.line_spacing = Pt(11)
d.styles['Ch5Code'].paragraph_format.widow_control = False
style('Ch5Table', 'Times New Roman', 10.5, after=0, line=1.05)
style('Ch5Caption', 'Times New Roman', 11, before=6, after=6, bold=True).paragraph_format.keep_with_next = True
chapter_title = style('Ch5Title', 'Times New Roman', 14, after=12, bold=True)
chapter_title.base_style = d.styles['Title']
chapter_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
chapter_title.paragraph_format.keep_with_next = True

def p(text='', bold=False):
    q = d.add_paragraph(text, 'Ch5Body')
    q.alignment = WD_ALIGN_PARAGRAPH.LEFT if '\n' in text else WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold:
        for run in q.runs:
            run.bold = True
    return q

def h(text, sub=False):
    return d.add_paragraph(text, 'Ch5Subheading' if sub else 'Ch5Heading')

def page(title, subtitle=None):
    if subtitle:
        first = d.add_paragraph(title, 'Ch5Title')
        d.add_paragraph(subtitle, 'Ch5Title')
    else:
        first = h(title)
    first.paragraph_format.page_break_before = True

def table(headers, rows, widths, caption=None):
    if caption:
        d.add_paragraph(caption, 'Ch5Caption')
    t = d.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = OxmlElement('w:' + side)
        for key, value in [('val','single'), ('sz','4'), ('color','D9D9D9')]:
            edge.set(qn('w:' + key), value)
        borders.append(edge)
    pr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, value in [('top','95'), ('bottom','95'), ('left','100'), ('right','100')]:
        el = OxmlElement('w:' + side)
        el.set(qn('w:w'), value)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    pr.append(margins)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for values in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = str(text)
    repeat = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(no_split)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7E6E6')
                cell._tc.get_or_add_tcPr().append(shade)
            for paragraph in cell.paragraphs:
                paragraph.style = d.styles['Ch5Table']
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                if i == 0:
                    for run in paragraph.runs:
                        run.bold = True
    d.add_paragraph('', 'Ch5Source')
    return t

listings = []
def excerpt(file, start, end, title, tail=0):
    lines = (ROOT / file).read_text(encoding='utf-8').splitlines()
    a = next(i for i, x in enumerate(lines) if start in x)
    b = next(i for i in range(a, len(lines)) if end in lines[i]) + tail
    body = textwrap.dedent('\n'.join(lines[a:b+1]))
    label = f'Listing {len(listings)+1}  {title}'
    d.add_paragraph(label, 'Ch5Caption')
    d.add_paragraph(f'Source: {file}\nLines {a+1} to {b+1}  Selected source excerpt', 'Ch5Source')
    last_code = None
    for line in body.splitlines():
        if line.strip():
            last_code = d.add_paragraph(line, 'Ch5Code')
    if last_code is not None:
        last_code.paragraph_format.space_after = Pt(8)
    listings.append({'listing':len(listings)+1, 'file':file, 'start':a+1, 'end':b+1, 'title':title, 'code':body})
    return body

# The source chapter consists of nine deliberately paginated pages.
page('Chapter 5', 'THE SOURCE CODE')
p('This chapter presents selected source code of the Inventory and Sales System for Paolo Paolo D.A. Matting and Accessories. The listings explain how the application controls access, maintains product and stock records, processes storefront sales, prepares reports, prints receipts, and supports password recovery and database restoration.')
p('The application uses PHP and the Laravel framework. Controllers process requests, models access database records, Blade templates display interfaces, and middleware enforces access restrictions. The current project includes SQL Server connection and snapshot support. Its dependency file requires PHP 8.3 or later within the supported version range and Laravel 13.8 or later within version 13.')
table(['Area', 'Main source files', 'Responsibility'], [
    ['Authentication', 'AuthenticatedSessionController.php\nEnsureUserRole.php', 'Login, session renewal and access control'],
    ['Inventory', 'Product.php\nProductController.php\nStockInController.php', 'Product details, stock thresholds and incoming batches'],
    ['Sales', 'PosController.php\nTransactionController.php', 'Checkout, stock deductions and sales history'],
    ['Reports', 'ReportController.php', 'Sales summaries and inventory information'],
    ['Recovery', 'AccountRecoveryService.php\nBackupController.php', 'Recovery codes and password-protected restore'],
    ['Receipt', 'pos/receipt.blade.php', 'Thermal receipt display and printing'],
], [1.03,2.65,2.32])
p('Each listing identifies its source file and line range. These are selected excerpts; declarations and surrounding statements remain in the project. Blank lines are removed for space. The listings cover product sales. Online orders, payroll and physical installation work are outside their scope.')

page('5.1 User Authentication and Access Control')
p('A login request authenticates the username and password before renewing the session identifier. Administrators are directed to the dashboard, while operational users are directed to the POS interface. The login request also checks account activity and verification status.')
excerpt('app/Http/Controllers/Auth/AuthenticatedSessionController.php', 'public function store(', "return redirect()->route('pos.index');", 'Authentication and role based redirection', 1)
p('Administrative modules are protected at the route level. Authentication and active-account middleware are applied before role restrictions, so hiding a menu alone does not grant or remove permission.')
excerpt('routes/web.php', "Route::middleware(['role:Admin'])->group", "->name('security.index');", 'Administrator route restriction')
p('The project also applies login lockouts after repeated failed attempts. Account disabling is enforced during authentication and subsequent protected requests.')

page('5.2 Product Records and Stock Thresholds')
p('Inventory quantities are calculated from remaining stock batches that are in good condition and have not expired. Damaged, defective, depleted and expired batches are excluded from sellable stock.')
excerpt('app/Models/Product.php', 'public function sellableStockIns()', "return (float) $this->sellableStockIns()->sum('Remaining_Quantity');", 'Sellable stock calculation', 1)
p('The low-stock threshold is stored with each product. Editing a product updates this saved value; when the field is omitted, the existing threshold is retained. Validation requires a non-negative integer.')
excerpt('app/Http/Controllers/ProductController.php', '$product->update([', "'Images' => $currentImages,", 'Saving an edited stock threshold', 1)

page('5.3 Recording Incoming Stock')
p('A stock-in transaction records the product, responsible user, received quantity, cost and retail prices, expiration information and condition. The received quantity becomes the batch remaining quantity when the record is created.')
excerpt('app/Http/Controllers/StockInController.php', "$hasExpiration = $request->boolean('Has_Expiration');", "'Condition' => $validated['Condition'] ?? 'Good',", 'Creating a stock batch', 1)
p('Before this code runs, the controller validates required fields, numeric quantities and prices, allowed product conditions, and any required expiration date. New products and their batches are created inside a database transaction. The product and batch records are therefore committed together.')
p('Physical inspection remains the responsibility of the owner or authorized employee. The application records the quantity and condition entered after inspection; it does not independently examine delivered products.')

page('5.4 Checking Available Stock and Allocating Batches')
p('Checkout locks eligible batches in ascending batch-ID order during the database transaction. First-in, first-out allocation follows batch creation order.')
excerpt('app/Http/Controllers/PosController.php', "$batches = StockIn::where('Product_ID', $product->ID)", "$available = (float) $batches->sum('Remaining_Quantity');", 'Selecting and locking sellable batches')
p('Requests exceeding sellable stock are rejected. Accepted quantities are allocated from the earliest eligible batch before continuing to later batches.')
excerpt('app/Http/Controllers/PosController.php', '$remainingToAllocate = $qty;', '$remainingToAllocate -= $take;', 'Allocating stock across batches', 2)

page('5.5 Saving Sales and Deducting Stock')
p('Checkout calculates totals on the server and validates payment before saving. Cash received must cover the total; change is the excess received. GCash payments require an exact matching amount and a 13-digit reference. Staff still verify that the payment was actually received in the GCash account.')
excerpt('app/Http/Controllers/PosController.php', '$sale = Sale::create([', '$batch->save();', 'Recording a sale and reducing batch balances', 2)
p('The sale header, sold items and stock deductions execute within the same database transaction. The completed transaction supplies its stored information to sales history, reporting and receipt output.')

page('5.6 Reports and Sales History')
p('The Reports module summarizes recorded sales and prepares inventory information. Sales summaries are computed from the same filtered query used to retrieve transaction records, keeping the displayed totals consistent with the selected scope and period.')
excerpt('app/Http/Controllers/ReportController.php', '$this->applySalePeriod($salesQuery,', "$gcashSales = (float) (clone $salesQuery)", 'Filtered sales totals', 2)
p('The report controller supports weekly, monthly, yearly and overall periods. Printing and PDF output use the generated report data.')
excerpt('app/Http/Controllers/TransactionController.php', 'private function buildQuery(', "// Apply Date / Period", 'Restricting employee sales history')
p('The administrator can review all transactions or filter by user. Non-administrative users are restricted to their own sales. Receipt access applies a corresponding ownership check.')

page('5.7 Password Recovery and Database Restoration')
p('Password recovery uses a private saved recovery code. The generated code contains 16 random characters arranged in four groups. Its hash is stored; the displayed code must be saved privately by its owner.')
excerpt('app/Services/AccountRecoveryService.php', 'public function generate(User $user): string', 'return $code;', 'Generating and hashing a recovery code', 2)
p('Resetting a password replaces the recovery code and removes the account\'s existing database sessions. Five failed attempts trigger a 15-minute lock. A reset grant lasts ten minutes.')
excerpt('app/Http/Controllers/BackupController.php', 'public function restoreBackup(', "'current_password.current_password' =>", 'Confirming the administrator password before restore', 1)
p('The administrator-only restore route checks the selected backup and uses the snapshot service for SQL Server or JSON restoration. Recovery codes must be regenerated after restoration.')

page('5.8 Thermal Receipt Output')
p('The receipt template displays stored transaction details, purchased items, payment information and totals. It supports 80 mm and 58 mm thermal paper. Printing hides the on-screen controls and removes screen-only margins and decoration.')
excerpt('resources/views/pos/receipt.blade.php', 'body[data-paper-width="58"] {', '--receipt-font-size: 8pt;', 'Selecting the smaller thermal paper layout', 1)
excerpt('resources/views/pos/receipt.blade.php', '@media print {', '.receipt-paper { margin: 0; box-shadow: none; }', 'Thermal printing styles', 1)
excerpt('resources/views/pos/receipt.blade.php', '<div class="row total bold">', "@endif", 'Displaying receipt payment totals')
p('The operator selects the matching printer paper size, uses 100 percent scale, removes margins, and disables browser headers and footers. The receipt has no logo and the action is labelled Print Receipt.')

# Appendix A: letter template, not a claim of granted permission.
letter_start = len(d.paragraphs)
page('APPENDIX A', 'LETTER TO THE COMPANY')
p('Date: ____________________')
p('Mr. and Mrs. Pabelonia\nOwners\nPaolo Paolo D.A. Matting and Accessories\nCompany address: ______________________________')
p('Subject: Request for Permission to Conduct a System Study', True)
p('Dear Mr. and Mrs. Pabelonia:')
p('We are students of the University of Mindanao enrolled in IT12/L, System Integration and Architecture. We respectfully request permission to conduct a study of the physical storefront inventory and sales processes of Paolo Paolo D.A. Matting and Accessories as part of our course requirements.')
p('Our project, Inventory and Sales System for Paolo Paolo D.A. Matting and Accessories, focuses on inventory monitoring, storefront sales recording, report preparation, user access, and database backup and recovery. With your permission, we would like to interview the owner or authorized staff, observe the relevant workflow, and review sample business forms needed to understand these processes.')
p('We will coordinate the schedule with you to avoid disrupting business activities. Information will be used for the academic study. Photographs and copies of records will be taken only with permission, and personal details, passwords, payment references and other sensitive information will be removed from materials included in the documentation.')
p('We would also appreciate your feedback during the demonstration and evaluation of the developed system. Thank you for considering our request.')
p('Respectfully yours,\n\nCharles Jayson R. Roque\nNiño Angelo C. Cirunay\nStudent Researchers')
p('Noted by: __________________________\nMeljohn V. Aborde\nIT12/L Subject Teacher')
p('Company acknowledgement\nName and signature: _______________________    Date: ____________')
for paragraph in d.paragraphs[letter_start:]:
    if paragraph.style.name == 'Ch5Body':
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(7)

# Appendix B: user explicitly requested examples.
page('APPENDIX B', 'COMPANY FORMS')
h('B.1 Sample Sales Transaction Record', True)
p('The following examples illustrate the fields used to document storefront transactions and incoming stock. Product names reflect the automotive items discussed in Chapter 1. All quantities, prices, dates and document numbers below are illustrative.')
p('Paolo Paolo D.A. Matting and Accessories\nSAMPLE SALES RECORD\nDocument No.: SAMPLE-001    Date: ____________________\nProcessed by: Authorized Cashier    Payment: Cash')
table(['Product description', 'Qty.', 'Unit price PHP', 'Amount PHP'], [
    ['Engine oil bottle', '2', '250.00', '500.00'],
    ['Oil filter', '1', '180.00', '180.00'],
    ['Wiper blade', '1', '220.00', '220.00'],
], [2.85,0.5,1.3,1.35])
p('Total amount: PHP 900.00\nAmount received: PHP 1,000.00\nChange: PHP 100.00\nGCash reference: Not applicable\nCustomer name or other details, when needed: ____________________')
p('For a GCash transaction, the amount received must match the total and the payment reference must contain 13 digits. Record the actual reference only after confirming the received payment. The example above is a cash transaction.')
p('This is an example of an operational transaction record for academic illustration. Replace it with a permitted copy of the company\'s actual form when available.')

page('APPENDIX B  COMPANY FORMS')
h('B.2 Sample Incoming Stock Record', True)
p('Stock-in No.: SAMPLE-SI-001    Date received: ____________________\nSupplier or purchase reference: ______________________________\nReceived and inspected by: _________________________________')
table(['Product', 'Qty. received', 'Unit cost PHP', 'Retail price PHP', 'Condition'], [
    ['Engine oil bottle', '12', '200.00', '250.00', 'Good'],
    ['Oil filter', '6', '140.00', '180.00', 'Good'],
    ['Wiper blade', '8', '170.00', '220.00', 'Good'],
], [2.0,0.85,0.95,1.15,1.05])
p('Unit of counting: Piece. Each engine oil bottle is one piece.\nTotal pieces received: 26\nPurchase cost total: PHP 4,600.00\nExpected retail value: PHP 5,840.00')
h('Receiving Inspection', True)
table(['Inspection point', 'Record or result'], [
    ['Quantity compared with order or supplier document', '_______________________________'],
    ['Brand, product type, size and vehicle compatibility', '_______________________________'],
    ['Packaging, seals, leaks and visible damage', '_______________________________'],
    ['Expiration date, if applicable', '_______________________________'],
    ['Discrepancy and action taken', '_______________________________'],
], [2.8,3.2])
p('Accepted by: ___________________    Checked by: ___________________\nThe stock-in record is entered only after actual quantity and condition have been checked. Illustrative values above do not represent a real delivery.')

page('APPENDIX B  COMPANY FORMS')
h('B.3 Sample Inventory Monitoring Record', True)
p('Stock verification date: ____________________\nPrepared by: ____________________\nCounting basis: Pieces of sellable product')
table(['Product', 'Recorded stock', 'Physical count', 'Difference', 'Threshold'], [
    ['Engine oil bottle', '10', '10', '0', '5'],
    ['Oil filter', '4', '4', '0', '5'],
    ['Wiper blade', '0', '0', '0', '3'],
], [2.0,1.0,1.0,0.9,1.1])
p('Example interpretation: Engine oil is above its threshold. Oil filters are low in stock because four pieces remain against a threshold of five. Wiper blades are out of stock. Thresholds are editable product settings and the values in this form are examples.')
h('Reconciliation and Follow-up', True)
p('Reason for any difference: ____________________________________\n_________________________________________________________\nCorrective action or replenishment note: _________________________\n_________________________________________________________\nApproved by: ____________________    Date: ____________________')
p('Physical counts verify the recorded balances. Damage, defects, expired items and unrecorded movement must be investigated before a correction is entered. This example supports manual verification; it is not an additional purchasing or supplier-management module.')

# Appendix C intentionally contains real-world evidence slots, never invented events.
page('APPENDIX C', 'DATA GATHERING EVIDENCE')
h('C.1 Interview and Observation Record', True)
p('Complete this record using the actual interview notes, photographs and permitted copies of company records. Record the activity details and attach the corresponding evidence.')
p('Activity date and time: _______________________________________\nLocation or interview method: __________________________________\nCompany representative and role: _______________________________\nResearcher or interviewer: ____________________________________\nPermission to record or photograph: ____________________________')
h('Interview Guide', True)
questions = [
    'Which automotive products are monitored, and what unit is used to count each item?',
    'How are incoming products compared with the order or supplier document?',
    'How are damaged, defective or expired products identified and recorded?',
    'How does the owner determine when a product needs replenishment?',
    'What fields are actually written in the sales document and transaction log?',
    'How are cash amounts and received GCash payments verified?',
    'When and how are sold quantities deducted from inventory records?',
    'Which sales and inventory reports are prepared, and how often?',
    'Which tasks require owner access and which are performed by the cashier?',
    'What problems occur when finding past transactions or preparing reports?',
]
for i, question in enumerate(questions, 1):
    q = d.add_paragraph(f'{i}. {question}', 'Ch5Body')
    q.paragraph_format.space_after = Pt(4)
    q.paragraph_format.line_spacing = 1.15
p('Responses or notes attachment: ________________________________\nConfirmed by the representative: _______________________________')

page('APPENDIX C  DATA GATHERING EVIDENCE')
h('C.2 Evidence Register', True)
table(['ID', 'Evidence to attach', 'Date and source', 'Attachment or permission'], [
    ['C-01', 'Interview notes confirmed by owner or authorized staff', '____________', '____________'],
    ['C-02', 'Observation notes of stock checking and receiving', '____________', '____________'],
    ['C-03', 'Permitted photograph of shelves or storage', '____________', '____________'],
    ['C-04', 'Redacted sample sales document or manual order log', '____________', '____________'],
    ['C-05', 'Redacted stock record or supplier document', '____________', '____________'],
    ['C-06', 'System demonstration or feedback record', '____________', '____________'],
], [0.45,2.55,1.05,1.95])
h('C.3 Evidence Caption and Attachment', True)
p('Evidence ID: ____________________\nCaption or activity: _________________________________________\nDate and location: __________________________________________\nWhat the evidence shows: ____________________________________\nSource and permission: ______________________________________')
p('\nAttach an actual permitted photograph or scanned record here.\n\n\n')
p('Attach evidence of the existing business process described in Chapter 1. Remove private information from photographs and copied records.')

# Appendix D: executable plan with open result records, not fabricated passes.
plan_start = len(d.paragraphs)
page('APPENDIX D', 'TEST PLAN')
h('D.1 Purpose and Test Environment', True)
p('The test plan verifies that the developed system supports the physical storefront requirements described in Chapters 2 and 3. Testing covers Inventory, Sales, Reports, Security, thermal receipt output, and Backup and Recovery.')
p('Use a separate test database and test accounts for Admin/Owner and Cashier. Prepare good, damaged, defective and expired batches; sample Cash and GCash transactions; a thermal printer with supported paper; and a disposable backup destination. Confirm that the selected SQL Server connection, PHP driver and local server are working before execution.')
h('D.2 Test Procedure and Responsibilities', True)
p('The researchers prepare the data and execute functional, integration and access-control tests. The owner or authorized representative performs the acceptance review. Record the actual result, date, tester and supporting screenshot or output for every case. Log failed cases as defects, correct them, and repeat the affected tests.')
p('Backup restoration tests must run against a disposable database. Take a verified baseline backup before the test, compare restored values with the baseline, and confirm that the test does not change live business records.')
h('D.3 Entry and Completion Criteria', True)
p('Testing starts when the application launches, test users can sign in, required tables exist, and sample records and backups are available. Testing is complete when all planned cases have results, critical defects are resolved and retested, totals and stock balances reconcile, access restrictions hold, and the owner or representative records acceptance.')
p('Acceptance criteria follow the evaluation form used in Chapter 4: Product Quality 25 percent, Functional Testing 25 percent, Data Security 20 percent, and User-Friendliness 30 percent, for a total of 100 percent. Record any new evaluation separately from the existing assessment.')
p('Test date: __________    Build or version: __________\nDatabase: __________    Tester: ____________________\nRecord column: enter Pass or Fail and the evidence reference after execution.')
for paragraph in d.paragraphs[plan_start:]:
    if paragraph.style.name == 'Ch5Body':
        paragraph.paragraph_format.line_spacing = 1.15

cases = [
 ('TP-01','Sign in with an active, verified Admin account.','Login succeeds and the dashboard opens.'),
 ('TP-02','Enter a wrong password three times; retry during the lockout.','Errors are displayed and the temporary login lockout is enforced.'),
 ('TP-03','Disable a test employee account, then attempt login and a protected request.','The disabled account cannot use protected functions.'),
 ('TP-04','As Cashier, open Inventory, Reports, Security and Backup URLs directly.','Administrative routes refuse access; POS remains available.'),
 ('TP-05','Add a product with required details; repeat with a missing required value.','Valid details save; invalid data shows a validation error.'),
 ('TP-06','Receive 12 good units of an existing product.','One batch is recorded and sellable stock increases by 12.'),
 ('TP-07','Edit the stock threshold to 7; reload and sign in again; try a negative value.','The saved value remains 7; negative thresholds are rejected.'),
 ('TP-08','Filter inventory by date, price and stock; sort names and values both ways.','Displayed rows match filters and ascending or descending order.'),
 ('TP-09','Sell goods totaling PHP 900 and receive PHP 1,000 cash.','Sale total is PHP 900; change is PHP 100; stock decreases correctly.'),
 ('TP-10','Attempt the same sale with PHP 800 cash.','Underpayment is rejected; no sale or stock deduction is saved.'),
 ('TP-11','Use GCash, the exact total and a 13-digit test reference.','Transaction saves the payment method and reference with zero change.'),
 ('TP-12','Submit GCash references with 12 digits, 14 digits or letters.','Each invalid reference is rejected without saving a sale.'),
 ('TP-13','Use a valid reference but a GCash amount different from the total.','The mismatched payment amount is rejected.'),
 ('TP-14','Request more than available sellable stock.','Insufficient-stock message appears; stock and sales remain unchanged.'),
 ('TP-15','Prepare two good batches with 3 and 5 units; sell 4 units.','FIFO uses all 3 units of the first batch and 1 of the next; 4 remain.'),
 ('TP-16','Prepare damaged, defective, expired and zero-quantity batches.','Those batches are excluded from sellable stock and checkout.'),
 ('TP-17','As Cashier, inspect history and request another user\'s receipt URL.','Only own transactions are listed; another user\'s receipt is refused.'),
 ('TP-18','Generate weekly, monthly, yearly and overall sales reports.','Included dates, counts and totals agree with the test sales records.'),
 ('TP-19','Generate the inventory report with stock and category filters.','Quantities, thresholds and report rows agree with eligible records.'),
 ('TP-20','Print or export reports and print a receipt at 58 mm and 80 mm.','Outputs match stored totals; receipt text fits and controls are hidden.'),
 ('TP-21','Reset a test password using its saved recovery code; retry the old code.','New password works; old code fails; a replacement code is issued.'),
 ('TP-22','Enter an incorrect recovery code five times; test an expired reset grant.','The 15-minute lock and ten-minute reset-grant expiry are enforced.'),
 ('TP-23','Create a backup to the disposable destination; inspect the resulting file.','A nonempty backup is saved and can be selected for restoration.'),
 ('TP-24','In the test database, try wrong restore password, invalid file, then a valid backup.','Invalid attempts fail; valid restore recovers baseline data and clears recovery codes.'),
 ('TP-25','With dependencies installed, disconnect internet, launch locally and perform a test sale.','Local login, inventory, sale and receipt functions remain usable.'),
]

for group, (a,b) in enumerate([(0,8),(8,16),(16,25)], 1):
    page('APPENDIX D  TEST PLAN')
    h(['D.4 Authentication and Inventory Tests','D.5 Sales and Payment Tests','D.6 Reports Recovery and Offline Tests'][group-1], True)
    table(['Case', 'Procedure or input', 'Expected result', 'Result and evidence'],
          [[case,procedure,expected,'________\n________'] for case,procedure,expected in cases[a:b]],
          [0.6,2.07,2.52,0.81])
    p('Execution date: __________    Tester: ____________________\nDefect or retest references: ______________________________')
    if group == 3:
        p('Acceptance decision: Accepted / Accepted with revisions / Not accepted\nRepresentative and signature: ____________________    Date: __________')

d.save(PART)

# Append only new XML, preserving every original package part except the body
# and styles. Existing chapter content, images, relationships and page furniture
# are retained verbatim. No settings or private configuration are embedded.
with ZipFile(SOURCE) as source_zip, ZipFile(PART) as part_zip:
    body_root = etree.fromstring(source_zip.read('word/document.xml'))
    new_root = etree.fromstring(part_zip.read('word/document.xml'))
    old_body = body_root.find('w:body',NS)
    new_body = new_root.find('w:body',NS)
    original_children = [etree.tostring(e) for e in old_body]
    contents = old_body.find('w:tbl',NS)
    existing_contents_rows = contents.findall('w:tr',NS)
    original_contents_rows = [etree.tostring(row) for row in existing_contents_rows]
    for text, number in [
        ('Chapter 5: The Source Code', '46'),
        ('Appendices', ''),
        ('Appendix A Letter to the Company', '55'),
        ('Appendix B Company Forms', '56'),
        ('Appendix C Data Gathering Evidence', '59'),
        ('Appendix D Test Plan', '61'),
    ]:
        copied_row = deepcopy(existing_contents_rows[-1])
        for cell, value in zip(copied_row.findall('w:tc',NS),[text,number]):
            nodes = cell.findall('.//w:t',NS)
            nodes[0].text = value
            for node in nodes[1:]:
                node.text = ''
        contents.append(copied_row)
    last_section = old_body.find('w:sectPr',NS)
    for child in new_body:
        if child.tag != qn('w:sectPr'):
            last_section.addprevious(deepcopy(child))
    styles_root = etree.fromstring(source_zip.read('word/styles.xml'))
    new_styles = etree.fromstring(part_zip.read('word/styles.xml'))
    for st in new_styles:
        if (st.get(qn('w:styleId')) or '').startswith('Ch5'):
            styles_root.append(deepcopy(st))
    replacements = {
        'word/document.xml': etree.tostring(body_root,xml_declaration=True,encoding='UTF-8',standalone=True),
        'word/styles.xml': etree.tostring(styles_root,xml_declaration=True,encoding='UTF-8',standalone=True),
    }
    with ZipFile(FINAL,'w',ZIP_DEFLATED) as final_zip:
        for info in source_zip.infolist():
            final_zip.writestr(deepcopy(info),replacements.get(info.filename,source_zip.read(info.filename)))

    with ZipFile(FINAL) as check_zip:
        changed = [name for name in source_zip.namelist() if source_zip.read(name) != check_zip.read(name)]
        assert set(changed) == {'word/document.xml','word/styles.xml'},changed
        checked_body = etree.fromstring(check_zip.read('word/document.xml')).find('w:body',NS)
        for i,value in enumerate(original_children[:-1]):
            if old_body[i] is not contents:
                assert etree.tostring(checked_body[i]) == value
        checked_contents = checked_body.find('w:tbl',NS)
        assert all(etree.tostring(checked_contents.findall('w:tr',NS)[i]) == value for i,value in enumerate(original_contents_rows))
        assert etree.tostring(checked_body[-1]) == original_children[-1]

(OUT/'_work/listings.json').write_text(json.dumps(listings,indent=2,ensure_ascii=False),encoding='utf-8')
(OUT/'_work/artifact.md').write_text(f'''# Chapter 5 and appendices editing contract
Reference: {SOURCE}
Reference SHA256: {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}
Reference Word render: qa-reference/reference.pdf, 51 physical pages
Authority: existing DOCX and source repository
Page system: US Letter portrait, final body section margins top/bottom/right 1 inch, left 1.5 inches.
Body: source-derived Times New Roman 12 point, justified, 1.5 spacing.
Chapter title: centered 14 point bold, matching existing chapter title hierarchy.
Added source code: Courier New 9.5 point with 11 point pitch; verbatim selected line ranges, shared indentation and blank lines removed only.
Added tables: 10.5 point, explicit column widths, light-gray borders, repeating shaded header, non-splitting rows.
Scope: append nine source chapter pages, Appendices A to D. Extend manual contents with verified new section page numbers. Earlier chapter text and reference numbering preserved.
All original package parts are preserve-only except document.xml (append body) and styles.xml (append Ch5 styles).
No interviews, signatures, company approvals or new test results invented. Forms are examples; evidence and execution fields remain fillable.
Rendering: packaged renderer diagnosed unavailable LibreOffice on Windows. Native Word exports read-only copies; pypdfium2 rasterizes locally.
Fidelity checks: original body XML prefix identical except added contents rows; original contents rows identical; preserve-only ZIP parts identical; compare rendered first 51 pages except the extended contents page; source chapter at most ten pages.
''',encoding='utf-8')
print(json.dumps({'output':str(FINAL),'chapter_pages_planned':9,'test_cases':len(cases),'listings':len(listings),'modified_package_parts':changed},indent=2))
